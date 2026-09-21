#!/usr/bin/env python3
"""
doc_convert.py - 统一转换层（文档 → Markdown + 页面图片资产）

════════════════════════════════════════════════════════════════════════
为什么需要这一层
════════════════════════════════════════════════════════════════════════

为每种格式各写一套"正则式"文本抽取，注定追不上版式复杂度：表格被打散成单词、
图表整块丢失、扫描件抽到空。**根因不是某个 bug，是架构选型。**

正解与 Claude / OpenAI 官方对 PDF 的既有做法一致 —— 每页渲染成位图，与文本层
**一起**交给模型。表格、图表、多栏排版这些本地抽取器永远追不上的东西，恰恰是
视觉模型最强的地方；而触发本次改造的"纯图片 PDF"根本不需要 OCR，渲染成页图
模型就能读。

成本口径（DeepSeek 实测口径，2026-09）：服务端会把每张图自动缩放到约
1300×1300 像素量级，**单图 token 封顶 1024**。所以图像 token 不是风险，
真正的约束是请求体大小（内联图 48 MiB）。本模块的缩放不省 token，
意义是省请求体与上传耗时。

════════════════════════════════════════════════════════════════════════
模块边界
════════════════════════════════════════════════════════════════════════

**叶子模块**：只依赖标准库 + 懒加载 `pymupdf` / `python-docx` / `openpyxl` /
`python-pptx`。**不 import `attachments` / 引擎模块** —— 降级链由调用方编排
（`attachments._convert_document` 对 PDF 仍有兜底；Office 走本模块的
`convert_office`，见 docs/frontend/15）。

对外三件事：

| 函数 | 用途 |
| --- | --- |
| `convert_pdf` | PDF → 文本层 + 页图资产（附件通道与工具通道共用） |
| `convert_office` | docx / xlsx / pptx → 文本 + 表格结构（无视觉版式，末尾如实声明） |
| `tool_cache_dir` | 工具读文档时页图的落盘目录（`<workdir>/.aigent/pages/<key>/`） |

契约（`convert_pdf`）：

    {
      "markdown": str,   # 每页 `--- 第 N 页 ---` + 文本层 + `<!--img:pN-->` 锚点
      "pages": int,      # 原文档总页数
      "images": [{"id","page","path","width","height"}],
      "tables": int,     # 识别到的表格数（find_tables 命中数，无框线表格会漏）
      "converter": "pymupdf",
      "warnings": [str], # "诚实失败"通道：解不出来/没渲染的东西显式记下来
    }

图片资产落在 `out_dir/<att_id>.pages/pN.jpg`。锚点 id（`pN`）**就是文件名去掉
后缀**，展开侧据此把 Markdown 切成交错的文本块与图片块；不依赖 meta.json，
meta 丢了也能恢复。
"""
from __future__ import annotations

import hashlib
import os
import re
import shutil
import time
from pathlib import Path

from logger import get_logger

log = get_logger("doc_convert")

# 图片资产目录后缀：`<att_id>.pages/`
ASSETS_DIR_SUFFIX = ".pages"

# 渲染页图时 zoom 的上限。有些 PDF 画布极小（如 200×100pt），按长边算出的
# zoom 会大到几百倍，渲染出天文数字像素 —— 必须夹住。
RENDER_ZOOM_CAP = 6.0
RENDER_MIN_ZOOM = 0.1

# 页图 JPEG 质量。比 prepare_image 的 q82 高一些：页面上的小字对 JPEG 伪影
# 远比照片敏感，而服务端反正会把图缩到 ~1300px，体积和清晰度都还有余量。
PAGE_JPEG_QUALITY = 88

# Markdown 里的图片锚点。`pN` = 页图，`fN` = 图内嵌要素（后续批次）。
IMAGE_ANCHOR_RE = re.compile(r"<!--\s*img:([A-Za-z][0-9]{1,4})\s*-->")

# ── Office 格式（2026-09-21 上移到本模块，见 docs/frontend/15）────────────
# 这三种格式**没有视觉版式**：LibreOffice 渲染（~800MB）已被否决，所以只做
# 「保序文本 + 表格结构」抽取，并在正文末尾**如实声明丢了什么**。声明写在
# 转换层而不是注入侧：附件通道与工具通道共用它，两处口径必须一致。
OFFICE_EXTS = (".docx", ".xlsx", ".pptx")

# 抽取结果为空时的占位（同时是**哨兵**：stage/UI 据此判定"已降级"而不是
# "解析成功但内容为空"）。前缀统一，各格式措辞不同。
EMPTY_NOTE_PREFIX = "（未提取到文本"

_OFFICE_LOSSY_NOTE = "（本格式仅提取文本与表格结构；图表、图片、版式未包含）"

# ── 工具读文档的页图缓存（2026-09-21，见 docs/frontend/15）──────────────
# 工具的页图必须**在工具调用期间**落盘：中性块只带路径，编码发生在发送边界，
# 而发送边界不许写盘。落在**工作空间内**的 `.aigent/pages/<key>/`，目录名由
# 「源文件绝对路径 + mtime_ns + size + 渲染参数」哈希得出 —— 源文件一变 key
# 就变，旧目录成为无人引用的死文件，因此**不需要精确 GC**，只需按 TTL 惰性清剪。
# `<workdir>/.aigent` 必须同时加进 `refs.DEFAULT_IGNORE_DIRS`，否则它会出现在
# `@` 的候选列表里（缓存目录对用户毫无引用价值）。
TOOL_CACHE_DIRNAME = ".aigent"
TOOL_CACHE_SUBDIR = "pages"
# 目录内自带的 .gitignore：即使用户的工作空间是 git 仓库，也不必改他们自己的
# .gitignore 去忽略我们的缓存 —— 我们自己把这个目录从他们的版本控制里摘出去。
_TOOL_CACHE_GITIGNORE = "*\n"


def asset_dir(out_dir: Path, att_id: str) -> Path:
    """该附件的图片资产目录。"""
    return Path(out_dir) / f"{att_id}{ASSETS_DIR_SUFFIX}"


def _page_has_picture(page) -> bool:
    """页面是否含内嵌位图（图表、照片、扫描页）。"""
    try:
        return bool(page.get_images(full=True))
    except Exception as exc:  # noqa: BLE001 - 探测失败不能拦住建页图
        log.debug("get_images 失败（按无图处理）: %s", exc)
        return False


def _page_has_drawings(page) -> bool:
    """页面是否有矢量绘制（线条/方框/曲线）。

    只用来区分"真空白页"与"纯矢量图形页"：没有文本、没有位图、没有表格，
    但画了一堆线条的页（示意图、流程图）**必须**渲染，否则内容就丢了。
    它**不参与**"文本够密就不渲染"的判断 —— 正文页常有页眉横线/表格框线，
    让矢量参与那个判断会导致几乎所有页都渲染。
    """
    try:
        return bool(page.get_drawings())
    except Exception as exc:  # noqa: BLE001
        log.debug("get_drawings 失败（按无矢量处理）: %s", exc)
        return False


def _count_tables(page) -> int:
    """`page.find_tables()` 命中数。**无框线、靠位置对齐的表格会漏**（实测命中 0），
    所以它只用来"多触发一次页图渲染"，不能当成表格数目的真相。"""
    try:
        finder = page.find_tables()
    except Exception as exc:  # noqa: BLE001 - 版本/API 差异不该毁掉整篇转换
        log.debug("find_tables 不可用（按 0 处理）: %s", exc)
        return 0
    return len(getattr(finder, "tables", []) or [])


def _needs_page_image(text: str, has_picture: bool, tables: int, min_text: int,
                      has_drawings: bool = False) -> bool:
    """这一页的**文本层是否足以代表页面内容**。

    文本层够密、又没有图/表 → 页图提供不了额外信息，渲染它纯属浪费请求体
    （请求体才是真正的成本约束，DeepSeek 内联上限 48 MiB）。反过来，出现下列
    任一信号就渲染：

    - 页面**完全没有文本**：只要还有位图/表格/矢量图形，就必须渲染；
      三者都没有才是真的空白页，跳过（不浪费一张图的额度）；
    - 文本层很短（低于 `min_text`）→ 大概率是扫描件或图为主的一页（事故形态）；
    - 有内嵌位图 → 有图表/照片；
    - find_tables 命中 → 有可识别的表格（行列关系正是文本层会丢掉的）。
    """
    if not text.strip():
        return bool(has_picture or tables or has_drawings)
    if len(text) < max(0, int(min_text)):
        return True
    return bool(has_picture or tables)


def _render_page(fitz, page, dest_dir: Path, page_no: int,
                 max_edge: int) -> dict | None:
    """整页渲染成 JPEG。返回资产描述，失败返回 None（调用方记 warning）。"""
    try:
        rect = page.rect
        long_edge = max(float(rect.width), float(rect.height)) or 1.0
        zoom = 1.0
        if max_edge and max_edge > 0:
            zoom = float(max_edge) / long_edge
        zoom = max(RENDER_MIN_ZOOM, min(zoom, RENDER_ZOOM_CAP))
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        name = f"p{page_no}.jpg"
        path = dest_dir / name
        # Pillow 不参与：PyMuPDF 自己就能落 JPEG，少一次解码/编码
        pix.save(str(path), jpg_quality=PAGE_JPEG_QUALITY)
        return {
            "id": f"p{page_no}",
            "page": page_no,
            "path": str(path),
            "width": int(pix.width),
            "height": int(pix.height),
        }
    except Exception as exc:  # noqa: BLE001 - 单页渲染失败不该毁掉整篇转换
        log.warning("第 %d 页渲染失败: %s: %s", page_no, type(exc).__name__, exc)
        return None


def convert_pdf(src: Path, out_dir: Path, att_id: str, *,
                max_edge: int = 1568, max_pages: int = 20, max_images: int = 20,
                min_text: int = 200) -> dict:
    """PDF → Markdown（含图片锚点）+ 页面图片资产。

    解析库缺失 / 文件损坏一律抛异常，由调用方（`attachments._convert_document`）
    回落纯文本抽取 —— 本模块不做静默降级，"抽不到内容"必须让用户看见。
    """
    try:
        import fitz  # pymupdf
    except ImportError as exc:  # pragma: no cover - 环境相关
        raise RuntimeError("未安装 pymupdf，无法转换 PDF") from exc
    # 官方静音开关：首次调用 find_tables() 会往 stdout 打一条
    # "Consider using the pymupdf_layout package…" 的推荐语，会污染应用日志。
    no_recommend = getattr(fitz, "no_recommend_layout", None)
    if callable(no_recommend):
        no_recommend()

    doc = fitz.open(str(src))
    try:
        total = len(doc)
        dest_dir = asset_dir(out_dir, att_id)
        parts: list[str] = []
        images: list[dict] = []
        warnings: list[str] = []
        tables_total = 0
        scan_pages: list[int] = []
        skipped_pages: list[int] = []
        # max_pages 只约束**页图数量**，文本层永远是全量的 —— 文本是检索与
        # 无视觉模型兜底的主通道，不该被图像预算砍掉。
        render_budget = total if max_pages <= 0 else min(total, max_pages)
        image_budget = max_images if max_images > 0 else total
        # 页图目录不可写时置位：后续页不再尝试渲染（见下面的 OSError 分支）
        render_disabled = False

        any_text = False
        for i in range(total):
            page = doc[i]
            page_no = i + 1
            text = (page.get_text() or "").strip()
            has_picture = _page_has_picture(page)
            tables = _count_tables(page)
            tables_total += tables

            parts.append(f"--- 第 {page_no} 页 ---")
            if text:
                parts.append(text)
                any_text = True

            if not _needs_page_image(text, has_picture, tables, min_text,
                                     _page_has_drawings(page)):
                continue
            if render_disabled:
                continue
            if page_no > render_budget or len(images) >= image_budget:
                skipped_pages.append(page_no)
                continue
            if not dest_dir.is_dir():
                try:
                    dest_dir.mkdir(parents=True, exist_ok=True)
                except OSError as exc:
                    # 页图目录可能是不可写的：工具通道把它落在**工作空间内**
                    # （用户可能挂在只读盘上 / 目录属主不对）。这种情况下
                    # "给不出页图"是可接受的降级 —— 但**文本层必须照旧完整**，
                    # 异常若穿出去会把文本一起丢掉（整轮对话被打死）。
                    log.warning("页图目录不可写，转为纯文本模式: %s", exc)
                    warnings.append("页图目录不可写，本次只提取文本层")
                    render_disabled = True
                    continue
            asset = _render_page(fitz, page, dest_dir, page_no, max_edge)
            if asset is None:
                warnings.append(f"第 {page_no} 页图像渲染失败")
                continue
            images.append(asset)
            parts.append(f"<!--img:{asset['id']}-->")
            if not text:
                # 只有**真的把这一页当图发出去**了才这么说。被跳过的空白页不该
                # 产生"已按图像发送"的告警 —— 那是假话。
                scan_pages.append(page_no)

        if scan_pages:
            preview = "、".join(str(p) for p in scan_pages[:5])
            more = f" 等 {len(scan_pages)} 页" if len(scan_pages) > 5 else ""
            warnings.append(f"第 {preview} 页无文本层{more}，已按图像发送")
        if skipped_pages:
            warnings.append(
                f"仅渲染前 {render_budget} 页图像（共 {total} 页），文本层保留全量")
        # 表格数不写进 warnings：`tables` 已单独成字段，写进 warnings 会让每个
        # 带表格的 PDF 都被 UI 标成"已降级"，而它其实完全正常。
        # 没有任何页图时清掉空目录，免得会话目录里留一堆空壳
        if dest_dir.is_dir() and not any(dest_dir.iterdir()):
            dest_dir.rmdir()

        markdown = "\n".join(parts)
        # 既没有文本、也没有页图 → 真的什么都没读到（如全空白 PDF）。交回调用方
        # 判定"未提取到文本"，而不是让一堆页码标记冒充内容。
        if not images and not any_text:
            markdown = ""

        return {
            "markdown": markdown,
            "pages": total,
            "images": images,
            "tables": tables_total,
            "converter": "pymupdf",
            "warnings": warnings,
        }
    finally:
        doc.close()


def find_anchors(markdown: str) -> list[str]:
    """Markdown 里出现过的锚点 id，按出现顺序（展开侧切块用）。"""
    return IMAGE_ANCHOR_RE.findall(markdown or "")


def humanize_anchors(markdown: str) -> str:
    """把存储锚点译成模型可读的说明。

    锚点（`<!--img:pN-->`）是**给展开侧切块用的存储标记**，不该出现在给模型看的
    文本里。`pN` 的 N 就是页码（见 `_render_page`），所以能还原成人话。
    """
    def _sub(match: re.Match) -> str:
        ident = match.group(1)
        digits = ident[1:]
        if ident.startswith("p") and digits.isdigit():
            return f"[第 {int(digits)} 页为图像，随附]"
        return f"[随附图片 {ident}]"

    return IMAGE_ANCHOR_RE.sub(_sub, markdown or "")


# ══════════════════════════════════════════════════════════════════
#  可调参数（读在调用点；见 AGENTS.md 的 config.json 约定）
# ══════════════════════════════════════════════════════════════════

_DEFAULT_TOOL_CACHE_TTL_SECONDS = 604800


def tool_cache_ttl_seconds() -> int:
    """页图缓存目录的存活秒数（默认 7 天）。超期目录在下一次渲染时被清剪。

    写成函数而不是常量：`config.load()` 合并配置进 `os.environ` 的时机可能晚于
    本模块被导入，常量会在配置生效前被固化。
    """
    raw = os.environ.get("TOOL_DOC_CACHE_TTL_SECONDS")
    if raw is None or not str(raw).strip():
        return _DEFAULT_TOOL_CACHE_TTL_SECONDS
    try:
        value = int(float(str(raw).strip()))
    except (TypeError, ValueError):
        return _DEFAULT_TOOL_CACHE_TTL_SECONDS
    return value if value > 0 else _DEFAULT_TOOL_CACHE_TTL_SECONDS


# ══════════════════════════════════════════════════════════════════
#  Office 文本抽取（docx / xlsx / pptx）
# ══════════════════════════════════════════════════════════════════
# 从 attachments.py 上移（2026-09-21，见 docs/frontend/15）：附件通道与工具读文档
# 通道**共用同一段抽取代码**，否则同一个 xlsx「上传」与「@ 引用」会给出不同质量
# 的结果。第三方库仍在**函数内**懒加载，本模块保持叶子性质（不 import attachments
# 或引擎模块）。
#
# 能力边界（明确记录，不是缺陷）：这三种格式**没有视觉版式** —— 只抽「保序文本 +
# 表格结构」，栏位、图文相对位置、图表一律拿不到。渲染它们需要 LibreOffice
# （~800MB 外部二进制），已被否决。所以正文末尾**如实声明**丢了什么，而不是让
# 模型以为已尽收眼底。

def empty_note(ext: str) -> str:
    """抽取结果为空时的占位串（各格式措辞不同，前缀统一）。

    它同时是**哨兵**：给模型的是"这里本该有内容但没读到"，给 stage/UI 的是
    「已降级」的判据（见 `text_is_empty_note`）。所以措辞不能随意改。
    """
    if ext == ".pdf":
        return "（未提取到文本，可能是扫描件或纯图片 PDF）"
    if ext == ".docx":
        return "（未提取到文本：文档可能只含图片、图表或文本框）"
    if ext == ".xlsx":
        return "（未提取到文本：工作簿可能只含图片或图表）"
    if ext == ".pptx":
        return "（未提取到文本：幻灯片可能只含图片）"
    return "（未提取到文本）"


def text_is_empty_note(text) -> bool:
    """正文是否只是「未提取到文本」占位（stage / UI 据此判定"已降级"）。"""
    return str(text or "").lstrip().startswith(EMPTY_NOTE_PREFIX)


def _extract_docx(path: Path) -> tuple[str, int]:
    """docx → (正文, 内容块数)。段落在前，表格按行 `" | "` 连接追加在后。"""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - 环境相关
        raise RuntimeError("未安装 python-docx，无法解析 .docx") from exc
    document = docx.Document(str(path))
    parts = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    blocks = len(parts)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
                blocks += 1
    return "\n".join(parts), blocks


def _extract_xlsx(path: Path) -> tuple[str, int]:
    """xlsx → (正文, 数据行数)。空工作表只留标题行、不计内容。"""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - 环境相关
        raise RuntimeError("未安装 openpyxl，无法解析 .xlsx") from exc
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        data_rows = 0
        for sheet in wb.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    rows.append("\t".join(cells))
            # 行数写进标题：模型据此知道这张表是「50 行」还是「5000 行」，
            # 也就知道下面的内容是不是被截断过。工作表标题**不算内容** —— 否则
            # 一个只有空表的工作簿也会抽出标题行，看起来"解析成功"实则什么都没读到。
            parts.append(f"--- 工作表: {sheet.title}（{len(rows)} 行）---")
            parts.extend(rows)
            data_rows += len(rows)
        return "\n".join(parts), data_rows
    finally:
        wb.close()


def _extract_pptx(path: Path) -> tuple[str, int]:
    """pptx → (正文, 有文本的文本框数)。纯图片幻灯片只留页标题、不计内容。"""
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - 环境相关
        raise RuntimeError("未安装 python-pptx，无法解析 .pptx") from exc
    prs = Presentation(str(path))
    parts: list[str] = []
    text_shapes = 0
    for idx, slide in enumerate(prs.slides, start=1):
        bodies: list[str] = []
        for shape in slide.shapes:
            frame = getattr(shape, "text_frame", None)
            if frame is not None and (frame.text or "").strip():
                bodies.append(frame.text.strip())
        # 同 xlsx：页标题不算内容
        parts.append(f"--- 第 {idx} 页 ---")
        parts.extend(bodies)
        text_shapes += len(bodies)
    return "\n".join(parts), text_shapes


_OFFICE_EXTRACTORS = {
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".pptx": _extract_pptx,
}

OFFICE_CONVERTER = "office_text"


def convert_office(src, ext: str, *, limit: int) -> dict:
    """docx / xlsx / pptx → Markdown（文本 + 表格结构）**并自行截断**。

    与 `convert_pdf` 的两处刻意差异：

    - **本函数自己截断**（`limit`）。PDF 的截断留给调用方，是因为页图与正文要按
      锚点交错、截断点必须由展开侧掌握；Office 正文是纯文本，截断与「已截断」的
      诚实声明必须一起走，否则两条通道（附件 / 工具）会各自实现一遍、迟早不一致。
    - 额外返回 `text_truncated`，调用方不必再比对长度。

    契约不变：库缺失 / 文件损坏一律抛异常，由调用方兜（不做静默降级）。
    """
    extractor = _OFFICE_EXTRACTORS.get(str(ext or "").lower())
    if extractor is None:
        raise ValueError(f"convert_office 不支持的格式：{ext}")
    body, units = extractor(Path(src))
    base = {
        "text_truncated": False,
        "pages": None,
        "images": [],
        "tables": 0,
        "converter": OFFICE_CONVERTER,
        "warnings": [],
    }
    if not units:
        # 抽空必须留痕：正文只有占位串，`text_is_empty_note` 是"已降级"的判据
        return {**base, "markdown": empty_note(ext)}
    truncated = len(body) > max(0, int(limit))
    if truncated:
        body = body[:limit]
    body = f"{body}\n\n{_OFFICE_LOSSY_NOTE}"
    if truncated:
        body += f"\n（已截断至 {limit} 字符，完整内容请用 bash 或脚本读取原文件）"
    warnings = ["内容已截断"] if truncated else []
    return {**base, "markdown": body, "text_truncated": truncated, "warnings": warnings}


# ══════════════════════════════════════════════════════════════════
#  工具读文档：页图缓存目录
# ══════════════════════════════════════════════════════════════════
# 为什么需要在工具调用期间落盘：中性图片块只带**路径**，真正的编码发生在发送
# 边界，而发送边界不许写盘（doc 14）。所以渲染必须发生在工具里，且产物要活到
# 那一跳请求为止 —— 落在 `<workdir>/.aigent/pages/<key>/`。
#
# 为什么不落 OS 临时目录：工具（`ToolRegistry`）只拿得到 `workdir` / `bash_cwd`，
# 拿不到会话元数据目录；而落在工作空间里意味着模型与用户都能用普通文件工具看到
# 它、清它。代价是与「引用通道零复制、不污染工作空间」出现一个明示例外 ——
# 用「@ 忽略清单 + 目录内 .gitignore + TTL 清剪」把它收窄，并在文档里如实写明。

def _cache_key(src, *, max_edge: int, max_pages: int, max_images: int,
               min_text: int) -> str | None:
    """缓存目录名 = 源文件身份 + 渲染参数的哈希。

    含 `mtime_ns` 与 `size`：文件一改，key 就变 —— 旧目录再无人引用，因此
    **不需要精确 GC**。渲染参数进 key 是因为它们决定页图集与渲染判据。
    """
    try:
        st = Path(src).resolve().stat()
    except (OSError, RuntimeError, ValueError):
        return None
    raw = "|".join((
        str(Path(src).resolve()), str(st.st_mtime_ns), str(st.st_size),
        str(max_edge), str(max_pages), str(max_images), str(min_text),
    ))
    return hashlib.sha1(raw.encode("utf-8", "replace")).hexdigest()[:16]


def _write_cache_gitignore(cache_root: Path) -> None:
    """建缓存根目录并写一份 `.gitignore`（内容 `*`）。

    不写 `.gitignore` 的话，用户的工作空间若是 git 仓库，一次读 PDF 就会在他们的
    `git status` 里冒出一堆未跟踪文件；而去改他们自己的 `.gitignore` 更越界。
    **目录创建也在这里做**（失败一律吞掉）：不可写是完全可接受的降级，
    由 `convert_pdf` 转为纯文本模式并记 warning。
    """
    try:
        cache_root.mkdir(parents=True, exist_ok=True)
        target = cache_root / ".gitignore"
        if not target.exists():
            target.write_text(_TOOL_CACHE_GITIGNORE, encoding="utf-8")
    except OSError as exc:  # noqa: BLE001 - 写不了就算了，不影响读文档
        log.debug("准备页图缓存目录失败（转为纯文本模式）: %s", exc)


def _prune_cache(pages_root: Path, keep: str) -> None:
    """清掉超过 TTL 的同级缓存目录（顺手做，不清剪也不影响正确性）。"""
    cutoff = time.time() - tool_cache_ttl_seconds()
    try:
        entries = list(os.scandir(pages_root))
    except OSError:
        return
    for entry in entries:
        if entry.name == keep:
            continue
        try:
            if not entry.is_dir(follow_symlinks=False):
                continue
            if entry.stat(follow_symlinks=False).st_mtime >= cutoff:
                continue
            shutil.rmtree(entry.path, ignore_errors=True)
        except OSError as exc:  # noqa: BLE001 - 单个目录删不掉不该毁掉本次读取
            log.debug("清剪缓存目录失败 %s: %s", entry.name, exc)


def tool_cache_dir(workdir, src, *, max_edge: int = 1568, max_pages: int = 20,
                   max_images: int = 20, min_text: int = 200) -> Path | None:
    """工具读文档时页图的落盘目录；无法**定位**时返回 None。

    **返回 None 只代表"连路径都算不出来"**（workdir 为空或不可解析）—— 那属于
    调用方的参数问题。目录**不可写**不在这里返回 None：那种情况是完全可接受的
    降级，由 `convert_pdf` 内部转为纯文本模式并记 warning。若这里因为不可写而返回
    None，调用方连**文本层**都拿不到，那才是真正的损失（用户可能只是把项目挂在
    只读盘上，读 PDF 的正文完全应当照常工作）。
    """
    if not str(workdir or "").strip():
        return None
    key = _cache_key(src, max_edge=max_edge, max_pages=max_pages,
                     max_images=max_images, min_text=min_text)
    if not key:
        return None
    try:
        cache_root = Path(workdir).expanduser().resolve() / TOOL_CACHE_DIRNAME
    except (OSError, RuntimeError, ValueError) as exc:
        log.warning("页图缓存路径不可解析: %s: %s", type(exc).__name__, exc)
        return None
    pages_root = cache_root / TOOL_CACHE_SUBDIR
    _write_cache_gitignore(cache_root)   # 内部建目录，失败静默
    _prune_cache(pages_root, key)
    return pages_root / key
