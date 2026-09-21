#!/usr/bin/env python3
"""统一转换层的回归测试 —— 2026-09-20。

覆盖 `agents/doc_convert.py`：PDF → Markdown（含图片锚点）+ 页面图片资产。

改造的靶心是"纯图片 PDF 抽不到内容"（本次事故），所以这里最重要的用例是
**扫描件必须渲染出页图**；同时守住反向的克制 —— 文本层已经够密的页不该白渲染
一张图（请求体是真正的成本约束，DeepSeek 内联上限 48 MiB）。

夹具用 pymupdf / Pillow 生成**真实可解析**的 PDF，不 mock 解析库。
运行：`.venv/bin/python -m unittest discover -s tests -v`（pytest 未安装，用内置 unittest）
"""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
AGENTS_DIR = ROOT / "agents"
if str(AGENTS_DIR) not in sys.path:
    sys.path.insert(0, str(AGENTS_DIR))

import doc_convert as D  # noqa: E402

# 用 ASCII 正文：pymupdf 的内置字体（Helvetica）无法编码中文，`insert_textbox`
# 会把中文写成 '?'，那样测的就不是"文本是否够密"而是字体替换了。
DENSE_TEXT = "This is a sufficiently long paragraph of body text. " * 12   # > 200 字符


class _PdfFixtureMixin:
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)
        self.src = self.root / "src"
        self.src.mkdir(parents=True, exist_ok=True)
        self.att_id = "att_test000001"

    def _png(self, name: str = "fig.png", size=(400, 300)) -> Path:
        from PIL import Image
        p = self.src / name
        Image.new("RGB", size, (200, 210, 230)).save(p)
        return p

    def _pdf_text(self, name: str = "text.pdf", text: str = DENSE_TEXT) -> Path:
        import fitz
        p = self.src / name
        doc = fitz.open()
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(40, 40, 560, 700), text, fontsize=10)
        doc.save(str(p))
        doc.close()
        return p

    def _pdf_image_only(self, name: str = "scan.pdf") -> Path:
        """只有图片、没有文本层 —— 触发本次事故的形态。"""
        import fitz
        p = self.src / name
        doc = fitz.open()
        page = doc.new_page()
        page.insert_image(fitz.Rect(40, 40, 400, 300), filename=str(self._png()))
        doc.save(str(p))
        doc.close()
        return p

    def _pdf_mixed(self, pages: int = 3, name: str = "mixed.pdf") -> Path:
        """每页交替：密集文本页 / 纯图片页。"""
        import fitz
        p = self.src / name
        doc = fitz.open()
        for i in range(pages):
            page = doc.new_page()
            if i % 2 == 0:
                page.insert_textbox(fitz.Rect(40, 40, 560, 700), DENSE_TEXT, fontsize=10)
            else:
                page.insert_image(fitz.Rect(40, 40, 400, 300), filename=str(self._png()))
        doc.save(str(p))
        doc.close()
        return p

    def _convert(self, path: Path, **kwargs) -> dict:
        return D.convert_pdf(path, self.root, self.att_id, **kwargs)

    def _assets(self) -> list[str]:
        d = D.asset_dir(self.root, self.att_id)
        return sorted(p.name for p in d.iterdir()) if d.is_dir() else []


class ConvertPdfTests(_PdfFixtureMixin, unittest.TestCase):
    def test_scanned_pdf_renders_page_image(self):
        """核心：无文本层的 PDF 必须产出页图 —— 这才是"内容给到模型"。"""
        out = self._convert(self._pdf_image_only())
        self.assertEqual(out["converter"], "pymupdf")
        self.assertEqual(out["pages"], 1)
        self.assertEqual(len(out["images"]), 1)
        self.assertEqual(self._assets(), ["p1.jpg"])
        self.assertIn("<!--img:p1-->", out["markdown"])
        self.assertTrue(any("无文本层" in w for w in out["warnings"]), out["warnings"])
        img = out["images"][0]
        self.assertEqual(img["id"], "p1")
        self.assertEqual(img["page"], 1)
        self.assertTrue(Path(img["path"]).is_file())
        self.assertGreater(img["width"], 0)
        self.assertGreater(img["height"], 0)

    def test_dense_text_page_is_not_rendered(self):
        """反向克制：文本层够密又没有图 → 页图提供不了额外信息，别白花请求体。"""
        out = self._convert(self._pdf_text())
        self.assertEqual(len(out["images"]), 0)
        self.assertEqual(self._assets(), [])
        self.assertNotIn("<!--img:", out["markdown"])
        self.assertIn(DENSE_TEXT[:20], out["markdown"])
        self.assertEqual(out["warnings"], [])

    def test_mixed_document_renders_only_pages_that_need_it(self):
        out = self._convert(self._pdf_mixed(pages=4))
        self.assertEqual(out["pages"], 4)
        # 页 2 / 4 是纯图片页；页 1 / 3 是密集文本页
        self.assertEqual([i["id"] for i in out["images"]], ["p2", "p4"])
        self.assertEqual(self._assets(), ["p2.jpg", "p4.jpg"])

    def test_anchor_position_matches_page_order(self):
        """锚点在 markdown 里的位置必须与页序一致 —— 展开侧靠它排交错块。"""
        out = self._convert(self._pdf_mixed(pages=4))
        self.assertEqual(D.find_anchors(out["markdown"]), ["p2", "p4"])
        first_page2 = out["markdown"].index("第 2 页")
        first_anchor = out["markdown"].index("<!--img:p2-->")
        page3 = out["markdown"].index("第 3 页")
        self.assertLess(first_page2, first_anchor)
        self.assertLess(first_anchor, page3)

    def test_max_images_caps_rendered_pages(self):
        out = self._convert(self._pdf_image_only("a.pdf"), max_images=1)
        self.assertEqual(len(out["images"]), 1)

    def test_max_pages_caps_rendered_pages_and_says_so(self):
        out = self._convert(self._pdf_mixed(pages=4), max_pages=2)
        self.assertEqual(len(out["images"]), 1)          # 只有页 2 需要渲染
        self.assertEqual(out["pages"], 4)
        self.assertTrue(any("仅渲染前" in w for w in out["warnings"]), out["warnings"])

    def test_text_layer_is_never_cut_by_page_budget(self):
        """页图预算不得削减文本层：文本是检索与无视觉模型兜底的主通道。"""
        out = self._convert(self._pdf_mixed(pages=4), max_pages=1, max_images=1)
        for page_no in (1, 2, 3, 4):
            self.assertIn(f"--- 第 {page_no} 页 ---", out["markdown"])

    def test_max_edge_bounds_rendered_size(self):
        out = self._convert(self._pdf_image_only(), max_edge=400)
        img = out["images"][0]
        self.assertLessEqual(max(img["width"], img["height"]), 420)   # 允许取整误差

    def test_rendered_size_is_scale_independent(self):
        """同一张图放在不同画布上，渲染出的长边都应落在 max_edge 附近 ——
        否则大画布 PDF（如 5625pt 的架构图）会渲出天文数字像素。"""
        import fitz
        big = self.src / "big_canvas.pdf"
        doc = fitz.open()
        page = doc.new_page(width=5625, height=3754)
        page.insert_image(fitz.Rect(200, 200, 800, 600), filename=str(self._png()))
        doc.save(str(big))
        doc.close()
        out = self._convert(big, max_edge=1568)
        img = out["images"][0]
        self.assertLessEqual(max(img["width"], img["height"]), 1600)

    def test_empty_pdf_yields_nothing(self):
        """全空白 PDF：不渲染页图、不留空目录，markdown 清空交回调用方判定。"""
        import fitz
        empty = self.src / "empty.pdf"
        doc = fitz.open()
        doc.new_page()
        doc.save(str(empty))
        doc.close()
        out = self._convert(empty)
        self.assertFalse(D.asset_dir(self.root, self.att_id).exists())
        self.assertEqual(out["images"], [])
        self.assertEqual(out["warnings"], [])
        # 页码标记不该冒充内容：没有文本也没有图 → 交回调用方说"未提取到文本"
        self.assertEqual(out["markdown"], "")

    def test_vector_only_page_is_rendered(self):
        """纯矢量示意图：没有文本、没有位图，但画了东西 → 必须渲染，否则内容全丢。"""
        import fitz
        p = self.src / "diagram.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.draw_rect(fitz.Rect(60, 60, 300, 200), color=(0, 0, 0), width=2)
        page.draw_line(fitz.Point(300, 130), fitz.Point(420, 130))
        doc.save(str(p))
        doc.close()
        out = self._convert(p)
        self.assertEqual(len(out["images"]), 1)
        self.assertEqual(self._assets(), ["p1.jpg"])
        self.assertTrue(any("无文本层" in w for w in out["warnings"]), out["warnings"])

    def test_missing_library_raises_so_caller_can_fall_back(self):
        """解析库缺失必须抛异常（而不是静默返回空）—— 降级链靠它触发。"""
        from unittest import mock
        with mock.patch.dict(sys.modules, {"fitz": None}):
            with self.assertRaises(Exception):
                self._convert(self._pdf_text())


class NeedsPageImageTests(unittest.TestCase):
    """渲染判据是纯函数，直接测边界。"""

    def test_short_text_needs_image(self):
        self.assertTrue(D._needs_page_image("很短的标题", False, 0, 200))

    def test_dense_text_without_visuals_does_not(self):
        self.assertFalse(D._needs_page_image("x" * 500, False, 0, 200))

    def test_embedded_picture_needs_image_regardless_of_text(self):
        self.assertTrue(D._needs_page_image("x" * 5000, True, 0, 200))

    def test_table_hit_needs_image_regardless_of_text(self):
        self.assertTrue(D._needs_page_image("x" * 5000, False, 2, 200))

    def test_blank_page_is_skipped(self):
        """真空白页（无文本/位图/表格/矢量）不渲染 —— 不浪费一张图的额度。"""
        self.assertFalse(D._needs_page_image("", False, 0, 200, False))

    def test_image_only_page_is_rendered(self):
        self.assertTrue(D._needs_page_image("", True, 0, 200, False))

    def test_vector_only_page_is_rendered(self):
        self.assertTrue(D._needs_page_image("", False, 0, 200, True))

    def test_threshold_zero_disables_the_text_length_rule(self):
        """min_text=0 → 只按"有图/有表"渲染，纯文本页一律不渲染。"""
        self.assertFalse(D._needs_page_image("x" * 10, False, 0, 0))


class AnchorTests(unittest.TestCase):
    def test_find_anchors_in_order(self):
        md = "a<!--img:p1-->b<!-- img:p2 -->c<!--img:f3-->"
        self.assertEqual(D.find_anchors(md), ["p1", "p2", "f3"])

    def test_non_anchor_comments_ignored(self):
        self.assertEqual(D.find_anchors("<!-- 普通注释 --><!--img: --><!--img:1-->"), [])

    def test_asset_dir_naming(self):
        self.assertEqual(D.asset_dir(Path("/x"), "att_abc").name, "att_abc.pages")


# ══════════════════════════════════════════════════════════════════
#  Office 抽取（2026-09-21 从 attachments 上移，见 docs/frontend/15）
# ══════════════════════════════════════════════════════════════════
# 上移的理由是**两条通道共用同一段代码**：附件通道与工具读文档通道对同一个
# xlsx 必须给出同样质量的结果。所以这组测试同时守两件事 —— 抽取内容本身，
# 以及末尾那条"丢了什么"的诚实声明（它以前根本不存在）。

class _OfficeFixture(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        # resolve()：macOS 的临时目录经 /var → /private/var 符号链接，不解析的话
        # 与 `tool_cache_dir` 内部的 resolve() 结果对不上（路径断言会假失败）
        self.root = Path(self._tmp.name).resolve()

    def _xlsx(self, name="book.xlsx", rows=(("月份", "金额"), ("一月", 12))) -> Path:
        import openpyxl
        p = self.root / name
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "销量"
        for row in rows:
            ws.append(list(row))
        wb.save(str(p))
        return p

    def _docx(self, name="doc.docx") -> Path:
        import docx
        p = self.root / name
        d = docx.Document()
        d.add_paragraph("第一段")
        d.add_paragraph("第二段")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "A"
        t.rows[0].cells[1].text = "B"
        d.save(str(p))
        return p

    def _pptx(self, name="deck.pptx") -> Path:
        from pptx import Presentation
        p = self.root / name
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "标题一"
        prs.save(str(p))
        return p


class ConvertOfficeTests(_OfficeFixture):
    def test_xlsx_carries_sheet_row_count_and_loss_note(self):
        out = D.convert_office(self._xlsx(), ".xlsx", limit=30000)
        self.assertEqual(out["converter"], D.OFFICE_CONVERTER)
        self.assertIn("--- 工作表: 销量（2 行）---", out["markdown"])
        self.assertIn("月份\t金额", out["markdown"])
        self.assertIn("仅提取文本与表格结构", out["markdown"])
        self.assertFalse(out["text_truncated"])
        self.assertEqual(out["images"], [])          # Office 永远没有页图
        self.assertEqual(out["warnings"], [])

    def test_docx_keeps_paragraphs_then_tables(self):
        out = D.convert_office(self._docx(), ".docx", limit=30000)
        body = out["markdown"]
        self.assertIn("第一段", body)
        self.assertIn("A | B", body)
        self.assertLess(body.index("第一段"), body.index("A | B"))

    def test_pptx_uses_slide_headers(self):
        out = D.convert_office(self._pptx(), ".pptx", limit=30000)
        self.assertIn("--- 第 1 页 ---", out["markdown"])
        self.assertIn("标题一", out["markdown"])

    def test_empty_workbook_returns_placeholder_without_loss_note(self):
        """抽空 → 只留占位串（`text_is_empty_note` 的判据），**不加**损失声明 ——
        声明是给"有内容但缺视觉"的正文用的，空文档要的是"什么都没读到"。"""
        out = D.convert_office(self._xlsx(rows=()), ".xlsx", limit=30000)
        self.assertTrue(D.text_is_empty_note(out["markdown"]))
        self.assertNotIn("仅提取文本与表格结构", out["markdown"])
        self.assertEqual(out["text_truncated"], False)

    def test_truncation_is_declared(self):
        out = D.convert_office(self._docx(), ".docx", limit=6)
        self.assertTrue(out["text_truncated"])
        self.assertEqual(out["warnings"], ["内容已截断"])
        self.assertIn("已截断至 6 字符", out["markdown"])
        # 损失声明必须**仍在**（截断说明追加在它之后，不能把它挤掉）
        self.assertIn("仅提取文本与表格结构", out["markdown"])

    def test_unsupported_extension_raises(self):
        with self.assertRaises(ValueError):
            D.convert_office(self.root / "a.pdf", ".pdf", limit=100)

    def test_missing_library_raises_so_caller_can_degrade(self):
        """库缺失必须抛（由调用方兜），不能静默返回空正文。"""
        import sys
        from unittest import mock
        book = self._xlsx()          # 先在补丁之外造好文件
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            with self.assertRaises(RuntimeError):
                D.convert_office(book, ".xlsx", limit=100)


class HumanizeAnchorTests(unittest.TestCase):
    def test_page_anchor_becomes_readable_note(self):
        md = "--- 第 1 页 ---\ntext\n<!--img:p3-->"
        self.assertEqual(D.humanize_anchors(md),
                         "--- 第 1 页 ---\ntext\n[第 3 页为图像，随附]")

    def test_unknown_anchor_kind_still_gets_a_label(self):
        self.assertEqual(D.humanize_anchors("x<!--img:f2-->y"),
                         "x[随附图片 f2]y")

    def test_no_anchor_is_untouched(self):
        self.assertEqual(D.humanize_anchors("plain text"), "plain text")


class ToolCacheDirTests(_OfficeFixture):
    def test_requires_a_workdir(self):
        """workdir 为空 → None（`Path("")` 会被解析成进程 cwd，那是另一个事故）。"""
        src = self._xlsx()
        self.assertIsNone(D.tool_cache_dir("", src))
        self.assertIsNone(D.tool_cache_dir(None, src))

    def test_dir_lives_under_the_workspace_and_is_keyed(self):
        src = self._xlsx()
        path = D.tool_cache_dir(self.root, src)
        self.assertIsNotNone(path)
        self.assertTrue(str(path).startswith(
            str(self.root / D.TOOL_CACHE_DIRNAME / D.TOOL_CACHE_SUBDIR)))
        self.assertEqual(path, D.tool_cache_dir(self.root, src))    # 稳定
        self.assertEqual(len(path.name), 16)                        # sha1 截断

    def test_gitignore_written_once(self):
        src = self._xlsx()
        D.tool_cache_dir(self.root, src)
        gi = self.root / D.TOOL_CACHE_DIRNAME / ".gitignore"
        self.assertEqual(gi.read_text(encoding="utf-8").strip(), "*")

    def test_absent_source_has_no_key(self):
        """源文件 stat 不到 → 算不出 key → 返回 None（调用方据此走错误分支）。"""
        ghost = self.root / "ghost.xlsx"
        self.assertIsNone(D._cache_key(ghost, max_edge=1, max_pages=1,
                                       max_images=1, min_text=1))
        self.assertIsNone(D.tool_cache_dir(self.root, ghost))


if __name__ == "__main__":
    unittest.main()
