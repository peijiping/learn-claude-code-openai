#!/usr/bin/env python3
"""附件存储 / 解析 / 展开的离线回归测试 —— 2026-09-20。

覆盖 `agents/attachments.py` 的全部对外契约（设计见
`docs/frontend/12-附件与文件输入.md`）：

  1. 登记        —— 复制而非移动（原文件必须在）、meta 落盘、拒绝非法输入
  2. 解析        —— pdf / docx / xlsx / pptx / 纯文本 / 图片 各自可抽到内容
  3. 归位        —— 草稿区 → 会话目录；幂等；att_id 找不到时标记 missing 而不报错
  4. 展开        —— 发送边界：图片 → image_url(data URL)、文档 → 文本块；
                    文件缺失 / 体积超限 / 块畸形 **一律降级不抛**
  5. 账本/回放   —— build_user_content / text_view / harvest_attachments
  6. 清理        —— 草稿 TTL、孤儿会话目录门限、只删目标会话、归档/还原不动

全部用例用临时目录承载（`WorkspacePaths` 注入），**不触碰真实 `~/.aigent`**。
运行：`.venv/bin/python -m unittest discover -s tests -v`（pytest 未安装，用内置 unittest）
"""
import json
import os
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
AGENTS_DIR = ROOT / "agents"
if str(AGENTS_DIR) not in sys.path:
    sys.path.insert(0, str(AGENTS_DIR))

import attachments as A  # noqa: E402
from paths import WorkspacePaths  # noqa: E402


def _workspace(root: Path) -> WorkspacePaths:
    """临时工作空间的路径束（data_root 与 workdir 都指向测试临时目录）。"""
    return WorkspacePaths("default", root, root, bash_cwd=root)


class _FixtureMixin:
    """生成真实可解析的样例文件（不 mock 解析库 —— 要测的就是它们真能read）。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)
        self.ws = _workspace(self.root)
        self.src = self.root / "src"
        self.src.mkdir(parents=True, exist_ok=True)
        A.clear_expand_cache()
        self.addCleanup(A.clear_expand_cache)

    def _file(self, name: str, text: str) -> Path:
        p = self.src / name
        p.write_text(text, encoding="utf-8")
        return p

    def _png(self, name: str = "shot.png", size=(12, 12)) -> Path:
        from PIL import Image
        p = self.src / name
        Image.new("RGB", size, (200, 30, 30)).save(p)
        return p

    def _pdf(self, name: str = "doc.pdf") -> Path:
        import fitz
        p = self.src / name
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Hello Attachment PDF")
        doc.save(str(p))
        doc.close()
        return p

    def _docx(self, name: str = "doc.docx") -> Path:
        import docx
        p = self.src / name
        d = docx.Document()
        d.add_paragraph("Word 正文第一段")
        d.add_paragraph("Word 正文第二段")
        d.save(str(p))
        return p

    def _xlsx(self, name: str = "book.xlsx") -> Path:
        import openpyxl
        p = self.src / name
        wb = openpyxl.Workbook()
        wb.active["A1"] = "列一"
        wb.active["B1"] = 42
        wb.save(str(p))
        return p

    def _pptx(self, name: str = "deck.pptx") -> Path:
        from pptx import Presentation
        p = self.src / name
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, 400, 100)
        box.text_frame.text = "幻灯片正文内容"
        prs.save(str(p))
        return p

    def _stage_ok(self, path: Path) -> dict:
        result = A.stage(self.ws, [str(path)])
        self.assertEqual(result["failed"], [], f"登记应当成功：{result['failed']}")
        self.assertEqual(len(result["items"]), 1)
        return result["items"][0]


# ══════════════════════════════════════════════════════════════════
class StageTests(_FixtureMixin, unittest.TestCase):
    """登记：复制而非移动 + 元数据 + 非法输入的可读拒绝。"""

    def test_stage_copies_file_and_leaves_original(self):
        src = self._file("note.txt", "正文内容")
        item = self._stage_ok(src)
        self.assertEqual(item["kind"], "text")
        self.assertEqual(item["name"], "note.txt")
        self.assertTrue(A._ATT_ID_RE.match(item["att_id"]))
        self.assertTrue(src.is_file(), "原文件必须原地留下（只复制，绝不 move）")
        draft = self.ws.attachments_dir / "_draft" / item["att_id"]
        self.assertTrue((draft / "meta.json").is_file())
        self.assertTrue((draft / f"{item['att_id']}.txt").is_file())

    def test_stage_text_file_does_not_duplicate_txt(self):
        """纯文本类**不**额外落 .txt（原件本身就是文本，展开时直接读它）。"""
        item = self._stage_ok(self._file("a.md", "# 标题"))
        draft = self.ws.attachments_dir / "_draft" / item["att_id"]
        self.assertEqual(item["kind"], "text")
        self.assertEqual(sorted(p.name for p in draft.iterdir()),
                         sorted(["meta.json", f"{item['att_id']}.md"]))

    def test_stage_pdf_extracts_pages(self):
        item = self._stage_ok(self._pdf())
        self.assertEqual(item["kind"], "document")
        self.assertEqual(item["pages"], 1)
        self.assertGreater(item["text_chars"], 0)

    def test_stage_office_formats_extract_text(self):
        for maker, expect in ((self._docx, "Word 正文第一段"),
                              (self._xlsx, "列一"),
                              (self._pptx, "幻灯片正文内容")):
            with self.subTest(maker=maker.__name__):
                item = self._stage_ok(maker())
                self.assertEqual(item["kind"], "document")
                draft = self.ws.attachments_dir / "_draft" / item["att_id"]
                body = (draft / f"{item['att_id']}.txt").read_text(encoding="utf-8")
                self.assertIn(expect, body)

    def test_stage_image_kind_and_no_text(self):
        item = self._stage_ok(self._png())
        self.assertEqual(item["kind"], "image")
        self.assertEqual(item["text_chars"], 0)

    def test_stage_image_resizes_only_oversized(self):
        small = self._stage_ok(self._png("small.png", (40, 40)))
        big = self._stage_ok(self._png("big.png", (2400, 300)))
        small_dir = self.ws.attachments_dir / "_draft" / small["att_id"]
        big_dir = self.ws.attachments_dir / "_draft" / big["att_id"]
        self.assertFalse((small_dir / f"{small['att_id']}.send.jpg").exists())
        self.assertTrue((big_dir / f"{big['att_id']}.send.jpg").exists())

    def test_stage_rejects_with_readable_reasons(self):
        cases = [
            (self.src / "nope.txt", "不存在"),
            (self.src, "文件夹"),
            (self._devnull(), "空文件"),
        ]
        for path, needle in cases:
            with self.subTest(path=str(path)):
                item, reason = A._stage_one(self.ws, self.ws.attachments_dir
                                            / "_draft", str(path))
                self.assertIsNone(item)
                self.assertIn(needle, reason)

    def _devnull(self) -> Path:
        p = self.src / "empty.txt"
        p.write_text("", encoding="utf-8")
        return p

    def test_stage_rejects_legacy_office_with_hint(self):
        p = self._file("old.doc", "x")
        result = A.stage(self.ws, [str(p)])
        self.assertEqual(result["items"], [])
        self.assertIn(".docx", result["failed"][0]["reason"])

    def test_stage_rejects_unknown_extension(self):
        result = A.stage(self.ws, [str(self._file("weird.zzz", "x"))])
        self.assertIn("不支持的文件类型", result["failed"][0]["reason"])

    def test_stage_rejects_oversized_with_limit_in_message(self):
        # 上限下限是 1KB（_int_env 的 minimum），故用 1024 作为收紧后的上限
        p = self.src / "big.txt"
        p.write_text("x" * 4000, encoding="utf-8")
        with mock.patch.dict(os.environ, {"ATTACHMENT_DOC_MAX_BYTES": "1024"}):
            result = A.stage(self.ws, [str(p)])
        self.assertEqual(result["items"], [])
        self.assertIn("过大", result["failed"][0]["reason"])

    def test_partial_failure_does_not_block_batch(self):
        """用户选 5 个文件，不能因为其中 1 个不支持就全失败。"""
        good = self._file("ok.txt", "fine")
        bad = self._file("bad.doc", "x")
        result = A.stage(self.ws, [str(bad), str(good)])
        self.assertEqual(len(result["items"]), 1)
        self.assertEqual(len(result["failed"]), 1)

    def test_duplicate_path_in_one_batch_staged_once(self):
        p = self._file("dup.txt", "x")
        result = A.stage(self.ws, [str(p), str(p)])
        self.assertEqual(len(result["items"]), 1)
        self.assertEqual(result["failed"], [])

    def test_unsupported_kind_never_raises(self):
        """桥层契约：stage 对任何输入都必须返回结构，不抛异常。

        空白路径直接跳过（不算失败），不存在的路径进 failed。
        """
        result = A.stage(self.ws, ["", "   ", "/no/such/path"])
        self.assertEqual(result["items"], [])
        self.assertEqual(len(result["failed"]), 1)
        self.assertEqual(result["failed"][0]["path"], "/no/such/path")


# ══════════════════════════════════════════════════════════════════
class MigrateTests(_FixtureMixin, unittest.TestCase):
    """归位：草稿区 → 会话目录，幂等，找不到时标记 missing。"""

    def _ref(self, item: dict) -> dict:
        return {"att_id": item["att_id"], "kind": item["kind"],
                "name": item["name"], "ext": item["ext"]}

    def test_migrate_moves_files_and_removes_draft(self):
        item = self._stage_ok(self._file("a.txt", "hi"))
        records = A.migrate_to_session(self.ws, [self._ref(item)], "SID0000001")
        self.assertEqual(len(records), 1)
        self.assertFalse(records[0]["missing"])
        session_dir = self.ws.attachments_dir / "SID0000001"
        self.assertTrue((session_dir / f"{item['att_id']}.txt").is_file())
        self.assertTrue((session_dir / f"{item['att_id']}.meta.json").is_file())
        self.assertFalse((self.ws.attachments_dir / "_draft" / item["att_id"]).exists())

    def test_migrate_is_idempotent(self):
        item = self._stage_ok(self._file("b.txt", "hi"))
        ref = self._ref(item)
        first = A.migrate_to_session(self.ws, [ref], "SID0000002")
        second = A.migrate_to_session(self.ws, [ref], "SID0000002")
        self.assertFalse(first[0]["missing"])
        self.assertFalse(second[0]["missing"])
        session_dir = self.ws.attachments_dir / "SID0000002"
        self.assertTrue((session_dir / f"{item['att_id']}.txt").is_file())

    def test_migrate_unknown_att_id_marks_missing_without_raising(self):
        records = A.migrate_to_session(
            self.ws, [{"att_id": "att_deadbeef00", "name": "x.txt",
                       "kind": "text", "ext": ".txt"}], "SID0000003")
        self.assertEqual(len(records), 1)
        self.assertTrue(records[0]["missing"])

    def test_migrate_ignores_malformed_refs(self):
        records = A.migrate_to_session(
            self.ws, ["junk", {"att_id": "../../etc/passwd"},
                      {"att_id": "xx"}], "SID0000004")
        self.assertEqual(records, [])

    def test_migrate_rejects_illegal_session_id(self):
        item = self._stage_ok(self._file("c.txt", "hi"))
        records = A.migrate_to_session(self.ws, [self._ref(item)], "../evil")
        self.assertEqual(records, [])


# ══════════════════════════════════════════════════════════════════
class ResolveTests(_FixtureMixin, unittest.TestCase):
    def test_resolve_finds_original_text_and_send_image(self):
        big = self._stage_ok(self._png("big.png", (2400, 300)))
        A.migrate_to_session(
            self.ws, [{"att_id": big["att_id"], "kind": "image",
                       "name": big["name"], "ext": big["ext"]}], "SID0000005")
        session_dir = self.ws.attachments_dir / "SID0000005"
        self.assertIsNotNone(A.resolve_files(session_dir, big)["original"])
        self.assertIsNotNone(A.resolve_files(session_dir, big)["send_image"])

    def test_resolve_document_text_file(self):
        item = self._stage_ok(self._pdf())
        A.migrate_to_session(
            self.ws, [{"att_id": item["att_id"], "kind": "document",
                       "name": item["name"], "ext": item["ext"]}], "SID0000006")
        session_dir = self.ws.attachments_dir / "SID0000006"
        files = A.resolve_files(session_dir, item)
        self.assertIsNotNone(files["text"])
        self.assertIsNotNone(files["original"])

    def test_resolve_returns_none_for_missing_dir(self):
        files = A.resolve_files(self.root / "nope", {"att_id": "att_abc123"})
        self.assertEqual(files, {"original": None, "text": None, "send_image": None})


# ══════════════════════════════════════════════════════════════════
class ExpandTests(_FixtureMixin, unittest.TestCase):
    """发送边界展开：契约是「不抛异常 + 无附件时原样返回同一对象」。

    这里刻意**走生产路径**取待发送消息（migrate 的真实记录 → build_user_content），
    而不是手搓附件块 —— 手搓会漏掉 `pages` / `text_truncated` 这类字段，
    测出来的"通过"和线上不是一回事。
    """

    def _prepare(self, item: dict, sid: str) -> tuple[dict, Path]:
        """归位 + 组装真实 user 消息，返回 `(消息, 会话目录)`。"""
        records = A.migrate_to_session(
            self.ws, [{"att_id": item["att_id"], "kind": item["kind"],
                       "name": item["name"], "ext": item["ext"]}], sid)
        msg = {"role": "user", "content": A.build_user_content("看看", records)}
        return msg, self.ws.attachments_dir / sid

    def _block(self, att_id: str, kind: str, name: str, ext: str) -> dict:
        """手工块：只用于「文件不在磁盘上」这类降级用例。"""
        return {"type": A.ATTACHMENT_BLOCK_TYPE, "attachment": {
            "id": att_id, "kind": kind, "name": name, "ext": ext}}

    def test_non_list_content_returns_same_object(self):
        msg = {"role": "user", "content": "纯文本"}
        self.assertIs(A.expand_content_for_model(msg, None), msg)

    def test_list_without_attachment_blocks_returns_same_object(self):
        msg = {"role": "user", "content": [{"type": "text", "text": "hi"}]}
        self.assertIs(A.expand_content_for_model(msg, None), msg)

    def test_malformed_attachment_block_passes_through_unchanged(self):
        msg = {"role": "user", "content": [
            {"type": A.ATTACHMENT_BLOCK_TYPE, "attachment": "not-a-dict"}]}
        self.assertIs(A.expand_content_for_model(msg, None), msg)

    def test_image_expands_to_data_url(self):
        item = self._stage_ok(self._png())
        msg, session_dir = self._prepare(item, "SID0000007")
        out = A.expand_content_for_model(msg, session_dir)
        block = out["content"][1]
        self.assertEqual(block["type"], "image_url")
        self.assertTrue(block["image_url"]["url"].startswith("data:image/png;base64,"))
        # 不回写原消息
        self.assertIsNot(out, msg)
        self.assertEqual(msg["content"][1]["type"], A.ATTACHMENT_BLOCK_TYPE)
        self.assertNotIn("base64", json.dumps(msg["content"]))

    def test_image_prefers_resized_copy_for_sending(self):
        item = self._stage_ok(self._png("big.png", (2400, 300)))
        msg, session_dir = self._prepare(item, "SID0000008")
        out = A.expand_content_for_model(msg, session_dir)
        self.assertTrue(
            out["content"][1]["image_url"]["url"].startswith("data:image/jpeg"))

    def test_missing_image_degrades_to_placeholder(self):
        msg = {"role": "user", "content": [
            self._block("att_missing001", "image", "gone.png", ".png")]}
        out = A.expand_content_for_model(msg, self.root / "nowhere")
        self.assertEqual(out["content"][0]["type"], "text")
        self.assertIn("缺失", out["content"][0]["text"])

    def test_document_expands_to_text_block_with_header(self):
        item = self._stage_ok(self._pdf())
        msg, session_dir = self._prepare(item, "SID0000009")
        block = A.expand_content_for_model(msg, session_dir)["content"][1]
        self.assertEqual(block["type"], "text")
        self.assertIn(f"[附件: {item['name']}]", block["text"])
        self.assertIn("Hello Attachment PDF", block["text"])
        self.assertIn("共 1 页", block["text"])

    def test_oversized_image_degrades_with_readable_note(self):
        # 造一张 >1KB 的图（下限 1KB 是 _int_env 的 minimum），再把内联上限收到 1KB
        from PIL import Image
        p = self.src / "noise.png"
        Image.frombytes("RGB", (60, 60), os.urandom(60 * 60 * 3)).save(p)
        item = self._stage_ok(p)
        msg, session_dir = self._prepare(item, "SID0000010")
        with mock.patch.dict(os.environ, {"ATTACHMENT_INLINE_MAX_BYTES": "1024"}):
            out = A.expand_content_for_model(msg, session_dir)
        self.assertEqual(out["content"][1]["type"], "text")
        self.assertIn("过大", out["content"][1]["text"])

    def test_expand_never_raises_on_broken_session_dir(self):
        """会话目录指向一个文件（不是目录）等异常输入也必须收成占位文本。"""
        weird = self.root / "weird"
        weird.write_text("not a dir", encoding="utf-8")
        out = A.expand_content_for_model(
            {"role": "user", "content": [
                self._block("att_x1234567", "document", "a.pdf", ".pdf")]}, weird)
        self.assertEqual(out["content"][0]["type"], "text")

    def test_data_url_cache_invalidates_on_change(self):
        item = self._stage_ok(self._png())
        msg, session_dir = self._prepare(item, "SID0000011")
        first = A.expand_content_for_model(msg, session_dir)["content"][1]
        second = A.expand_content_for_model(msg, session_dir)["content"][1]
        self.assertEqual(first, second)
        # 换掉文件内容（mtime 变化）→ 缓存必须失效
        from PIL import Image
        target = session_dir / f"{item['att_id']}.png"
        Image.new("RGB", (12, 12), (10, 200, 10)).save(target)
        os.utime(target, (time.time() + 2,) * 2)
        third = A.expand_content_for_model(msg, session_dir)["content"][1]
        self.assertNotEqual(third["image_url"]["url"], first["image_url"]["url"])


# ══════════════════════════════════════════════════════════════════
class LedgerShapeTests(_FixtureMixin, unittest.TestCase):
    """账本形态：build_user_content / text_view / harvest_attachments。"""

    def test_build_user_content_returns_str_without_attachments(self):
        """这是零行为变化的硬保证：无附件必须是 str，且就是原文。"""
        self.assertEqual(A.build_user_content("hi", []), "hi")
        self.assertIsInstance(A.build_user_content("", []), str)

    def test_build_user_content_shape(self):
        records = [{"att_id": "att_abcd1234", "kind": "image", "name": "a.png",
                    "mime": "image/png", "ext": ".png", "size": 3}]
        blocks = A.build_user_content("看看", records)
        self.assertIsInstance(blocks, list)
        self.assertEqual(blocks[0], {"type": "text", "text": "看看"})
        self.assertEqual(blocks[1]["type"], A.ATTACHMENT_BLOCK_TYPE)
        att = blocks[1]["attachment"]
        self.assertEqual(att["id"], "att_abcd1234")
        self.assertEqual(att["name"], "a.png")
        self.assertNotIn("data:", json.dumps(blocks), "账本里绝不能出现文件字节")

    def test_build_user_content_omits_empty_text_block(self):
        blocks = A.build_user_content("   ", [{"att_id": "att_abcd1234"}])
        self.assertEqual(len(blocks), 1)
        self.assertEqual(blocks[0]["type"], A.ATTACHMENT_BLOCK_TYPE)

    def test_harvest_roundtrip(self):
        records = [{"att_id": "att_abcd1234", "kind": "image", "name": "a.png",
                    "mime": "image/png", "ext": ".png", "size": 7,
                    "source_path": "/x/a.png",
                    "stored_path": "/y/.attachments/S1/att_abcd1234.png"}]
        harvested = A.harvest_attachments(A.build_user_content("x", records))
        self.assertEqual(harvested[0]["id"], "att_abcd1234")
        self.assertEqual(harvested[0]["name"], "a.png")
        self.assertEqual(harvested[0]["size"], 7)
        self.assertEqual(harvested[0]["stored_path"],
                         "/y/.attachments/S1/att_abcd1234.png")

    def test_harvest_skips_non_attachment_content(self):
        self.assertEqual(A.harvest_attachments("plain"), [])
        self.assertEqual(A.harvest_attachments([{"type": "text", "text": "x"}]), [])
        self.assertEqual(
            A.harvest_attachments([{"type": "attachment", "attachment": {}}]), [])

    def test_text_view_reads_attachments_without_bytes(self):
        content = A.build_user_content("看看这个", [
            {"att_id": "att_abcd1234", "kind": "image", "name": "a.png"}])
        text = A.text_view(content)
        self.assertIn("看看这个", text)
        self.assertIn("[图片: a.png]", text)
        self.assertNotIn("base64", text)
        self.assertEqual(A.text_view("plain"), "plain")


# ══════════════════════════════════════════════════════════════════
class GarbageCollectionTests(_FixtureMixin, unittest.TestCase):
    def _age(self, path: Path, seconds: float) -> None:
        old = time.time() - seconds
        os.utime(path, (old, old))

    def test_gc_drafts_removes_only_expired(self):
        old_dir = self.ws.attachments_dir / "_draft" / "att_old0000001"
        new_dir = self.ws.attachments_dir / "_draft" / "att_new0000001"
        old_dir.mkdir(parents=True)
        new_dir.mkdir(parents=True)
        (old_dir / "meta.json").write_text("{}", encoding="utf-8")
        (new_dir / "meta.json").write_text("{}", encoding="utf-8")
        self._age(old_dir, 4 * 24 * 3600)
        self.assertEqual(A.gc_drafts(self.ws), 1)
        self.assertFalse(old_dir.exists())
        self.assertTrue(new_dir.exists())

    def test_gc_orphan_removes_dir_without_session(self):
        orphan = self.ws.attachments_dir / "GONE000001"
        orphan.mkdir(parents=True)
        self._age(orphan, 3 * 24 * 3600)
        self.assertEqual(A.gc_orphan_session_dirs(self.ws), 1)
        self.assertFalse(orphan.exists())

    def test_gc_orphan_keeps_trashed_session_attachments(self):
        """归档（trash）只改 meta 状态、jsonl 原样保留 —— 附件必须留着。"""
        sid = "TRASHED001"
        (self.ws.chat_history_dir).mkdir(parents=True, exist_ok=True)
        (self.ws.chat_history_dir / f"session_{sid}.jsonl").write_text(
            '{"role":"user","content":"x"}\n', encoding="utf-8")
        (self.ws.chat_history_dir / f"session_{sid}.meta.json").write_text(
            '{"status":"trashed"}', encoding="utf-8")
        att_dir = self.ws.attachments_dir / sid
        att_dir.mkdir(parents=True)
        self._age(att_dir, 5 * 24 * 3600)
        self.assertEqual(A.gc_orphan_session_dirs(self.ws), 0)
        self.assertTrue(att_dir.exists())

    def test_gc_orphan_respects_min_age(self):
        fresh = self.ws.attachments_dir / "NEWONEOF01"
        fresh.mkdir(parents=True)
        self.assertEqual(A.gc_orphan_session_dirs(self.ws), 0)
        self.assertTrue(fresh.exists())

    def test_gc_orphan_never_touches_draft_dir(self):
        draft = self.ws.attachments_dir / "_draft"
        draft.mkdir(parents=True)
        self._age(draft, 30 * 24 * 3600)
        A.gc_orphan_session_dirs(self.ws)
        self.assertTrue(draft.exists())

    def test_remove_session_attachments_targets_one_session_only(self):
        keep = self.ws.attachments_dir / "KEEP000001"
        drop = self.ws.attachments_dir / "DROP000001"
        keep.mkdir(parents=True)
        drop.mkdir(parents=True)
        A.remove_session_attachments(self.ws, "DROP000001")
        self.assertFalse(drop.exists())
        self.assertTrue(keep.exists())

    def test_remove_session_attachments_ignores_illegal_id(self):
        A.remove_session_attachments(self.ws, "../..")  # 不应抛，也不应删任何东西
        self.assertTrue(self.ws.attachments_dir.parent.exists())


if __name__ == "__main__":
    unittest.main()
