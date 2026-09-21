#!/usr/bin/env python3
"""`run_read` 统一读取入口的离线回归测试 —— 2026-09-21（设计见 docs/frontend/15）。

被守护的东西：

  1. 分派    —— 文本 / 图片 / PDF / Office / 目录，按**魔数优先、扩展名兜底**路由
  2. PDF     —— 文本层 + 页图（含图表的页才有页图）；模型可见文本里**没有存储锚点**
  3. 页图缓存 —— 落在 `<workdir>/.aigent/pages/<key>/`；key 随源文件变化；自带 .gitignore
  4. 降级    —— 工作空间不可写 → 纯文本 + 警告（**绝不**连正文一起丢）
  5. 绝不抛  —— 转换层抛异常 → 收束成 `"Error: ..."` 字符串
  6. 聚合    —— 一次文档读取是**一个组**：要么整组随附，要么整组丢掉（不切页）
  7. refs    —— 注入文本不再需要格式→工具的映射表；`.aigent` 进忽略清单

**为什么这些用真文件**：判分派靠魔数、判页图靠"这一页的文本层够不够密"，
两者都是真实字节的属性。用假数据只能测出我们自己的假设，测不出 PDF 的真实行为。

运行：`.venv/bin/python -m unittest discover -s tests -v`
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
AGENTS_DIR = ROOT / "agents"
if str(AGENTS_DIR) not in sys.path:
    sys.path.insert(0, str(AGENTS_DIR))

import attachments as A  # noqa: E402
import doc_convert as D  # noqa: E402
import refs as R  # noqa: E402
from tools import ToolRegistry  # noqa: E402

# 长到足以让 `_needs_page_image` 判定"文本层足够代表这一页"（阈值默认 200）
DENSE_TEXT = ("The quick brown fox jumps over the lazy dog. " * 12).strip()


class _WorkdirCase(unittest.TestCase):
    """临时工作空间（**不触碰真实 ~/.aigent**）。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name).resolve()
        self.reg = ToolRegistry(workdir=self.root, bash_cwd=self.root)
        A.clear_expand_cache()
        self.addCleanup(A.clear_expand_cache)

    # ── 夹具 ───────────────────────────────────────────────────────
    def png(self, name="pic.png") -> Path:
        from PIL import Image
        path = self.root / name
        Image.new("RGB", (40, 30), (10, 120, 200)).save(path)
        return path

    def pdf(self, name="dense.pdf", pages=("text",)) -> Path:
        """按页类型表造 PDF：`text` = 密集文本页，`image` = 纯图片页。"""
        import fitz
        path = self.root / name
        doc = fitz.open()
        for kind in pages:
            page = doc.new_page()
            if kind == "text":
                page.insert_textbox(fitz.Rect(40, 40, 560, 760), DENSE_TEXT,
                                    fontsize=10)
            else:
                page.insert_image(fitz.Rect(40, 40, 300, 220),
                                  filename=str(self.png("_seed.png")))
        doc.save(str(path))
        doc.close()
        return path

    def xlsx(self, name="book.xlsx") -> Path:
        import openpyxl
        path = self.root / name
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "销量"
        ws.append(["月份", "金额"])
        ws.append(["一月", 12])
        wb.save(str(path))
        return path

    def docx(self, name="doc.docx") -> Path:
        import docx
        path = self.root / name
        d = docx.Document()
        d.add_paragraph("第一段正文")
        table = d.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "A"
        table.rows[0].cells[1].text = "B"
        d.save(str(path))
        return path

    def pptx(self, name="deck.pptx") -> Path:
        from pptx import Presentation
        path = self.root / name
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "标题一"
        prs.save(str(path))
        return path


# ══════════════════════════════════════════════════════════════════
#  1. 分派
# ══════════════════════════════════════════════════════════════════

class DispatchTests(_WorkdirCase):
    def test_text_file_returns_text(self):
        (self.root / "a.py").write_text("print(1)\n", encoding="utf-8")
        self.assertEqual(self.reg.run_read("a.py"), "print(1)")

    def test_text_limit_keeps_tail_hint(self):
        (self.root / "long.txt").write_text("\n".join(f"l{i}" for i in range(9)),
                                            encoding="utf-8")
        out = self.reg.run_read("long.txt", limit=3)
        self.assertEqual(out.splitlines()[:3], ["l0", "l1", "l2"])
        self.assertIn("6 more lines", out)

    def test_image_returns_neutral_block(self):
        self.png()
        out = self.reg.run_read("pic.png")
        self.assertTrue(A.is_tool_image_result(out), out)

    def test_pdf_dense_text_has_no_page_image(self):
        """文本层足够密 → 不渲染页图（渲染它纯属浪费请求体），返回纯文本。"""
        self.pdf("dense.pdf", pages=("text",))
        out = self.reg.run_read("dense.pdf")
        self.assertIsInstance(out, str)
        self.assertIn("共 1 页", out)
        self.assertIn("本次未随附页图", out)
        self.assertIn("quick brown fox", out)

    def test_pdf_scanned_page_comes_back_as_image(self):
        """纯图片页（无文本层）→ 页图随附 + 正文里标明是第几页。

        这条正是引用通道改造的靶心：改造前它只回一句"(无可提取文本，可能为扫描件)"。
        """
        self.pdf("scan.pdf", pages=("image",))
        out = self.reg.run_read("scan.pdf")
        self.assertTrue(A.is_tool_image_result(out), out)
        images = A.tool_image_items(out)
        self.assertEqual(len(images), 1)
        self.assertEqual(images[0]["page"], 1)
        self.assertEqual(images[0]["mime"], "image/jpeg")
        self.assertTrue(Path(images[0]["path"]).is_file())
        self.assertIn("随附页图 1 张（第 1 页）", out["text"])

    def test_pdf_model_visible_text_has_no_storage_anchor(self):
        """`<!--img:pN-->` 是**存储锚点**，不许出现在给模型看的文本里。"""
        self.pdf("mixed.pdf", pages=("text", "image"))
        out = self.reg.run_read("mixed.pdf")
        text = A.tool_image_text(out)
        self.assertNotIn("<!--img:", text)
        self.assertIn("[第 2 页为图像，随附]", text)
        # 第 1 页是密集文本页 → 不渲染页图，所以只有 1 张
        self.assertEqual(len(A.tool_image_items(out)), 1)

    def test_office_xlsx_has_sheet_header_and_loss_note(self):
        self.xlsx()
        out = self.reg.run_read("book.xlsx")
        self.assertIn("--- 工作表: 销量（2 行）---", out)
        self.assertIn("月份\t金额", out)
        # 诚实声明：图表 / 图片 / 版式没有包含进来
        self.assertIn("仅提取文本与表格结构", out)

    def test_office_docx_keeps_paragraphs_and_tables(self):
        self.docx()
        out = self.reg.run_read("doc.docx")
        self.assertIn("第一段正文", out)
        self.assertIn("A | B", out)
        self.assertIn("仅提取文本与表格结构", out)

    def test_office_pptx_page_headers(self):
        self.pptx()
        out = self.reg.run_read("deck.pptx")
        self.assertIn("--- 第 1 页 ---", out)
        self.assertIn("标题一", out)

    def test_directory_pointer_error(self):
        (self.root / "sub").mkdir()
        out = self.reg.run_read("sub")
        self.assertTrue(out.startswith("Error:"), out)
        self.assertIn("目录", out)

    def test_missing_file_and_escape_are_error_strings(self):
        for bad in ("nope.txt", "../outside.txt", "/etc/hosts"):
            out = self.reg.run_read(bad)
            self.assertIsInstance(out, str, bad)
            self.assertTrue(out.startswith("Error:"), f"{bad} -> {out}")

    def test_unknown_binary_hint(self):
        (self.root / "blob.bin").write_bytes(b"\x00\x01\x02\xff\xfe" * 40)
        out = self.reg.run_read("blob.bin")
        self.assertTrue(out.startswith("Error:"), out)
        self.assertIn("二进制", out)
        self.assertIn("file", out)

    def test_conversion_layer_exception_becomes_error_string(self):
        """契约：转换层抛异常必须被收束成字符串，不能穿透打死整轮。"""
        self.pdf("dense.pdf", pages=("text",))
        with mock.patch.object(D, "convert_pdf",
                               side_effect=RuntimeError("boom")):
            out = self.reg.run_read("dense.pdf")
        self.assertTrue(out.startswith("Error:"), out)
        self.assertIn("boom", out)

    def test_office_exception_becomes_error_string(self):
        self.xlsx()
        with mock.patch.object(D, "convert_office",
                               side_effect=RuntimeError("no openpyxl")):
            out = self.reg.run_read("book.xlsx")
        self.assertTrue(out.startswith("Error:"), out)
        self.assertIn("no openpyxl", out)

    def test_no_tool_name_is_unknown(self):
        """退役的两个名字必须真的从工具表与 handler 表里消失。"""
        names = {t["function"]["name"] for t in self.reg.base_tools}
        self.assertNotIn("run_read_pdf", names)
        self.assertNotIn("view_image", names)
        self.assertIsNone(self.reg.resolve_handler("run_read_pdf"))
        self.assertIsNone(self.reg.resolve_handler("view_image"))
        # 未知工具名仍然只得到一句错误，不抛
        self.assertTrue(self.reg.execute("view_image", path="x.png")
                        .startswith("Error: Unknown tool"))


# ══════════════════════════════════════════════════════════════════
#  2-5. 页图缓存与降级
# ══════════════════════════════════════════════════════════════════

class PageImageCacheTests(_WorkdirCase):
    def _cache_root(self) -> Path:
        return self.root / D.TOOL_CACHE_DIRNAME / D.TOOL_CACHE_SUBDIR

    def test_page_images_land_under_workdir_dot_aigent(self):
        pdf = self.pdf("scan.pdf", pages=("image",))
        out = self.reg.run_read("scan.pdf")
        path = Path(A.tool_image_items(out)[0]["path"])
        self.assertTrue(path.is_file())
        self.assertTrue(str(path).startswith(str(self._cache_root())),
                        f"{path} 不在 {self._cache_root()} 下")

    def test_cache_carries_its_own_gitignore(self):
        """工作空间是 git 仓库时，缓存不该污染用户的 `git status`。"""
        self.pdf("scan.pdf", pages=("image",))
        self.reg.run_read("scan.pdf")
        gi = self.root / D.TOOL_CACHE_DIRNAME / ".gitignore"
        self.assertTrue(gi.is_file())
        self.assertEqual(gi.read_text(encoding="utf-8").strip(), "*")

    def test_cache_key_follows_source_identity(self):
        """key 含 mtime/size → 源文件一改就换目录，旧目录成为死文件（无需精确 GC）。"""
        pdf = self.pdf("scan.pdf", pages=("image",))
        first = D.tool_cache_dir(self.root, pdf)
        self.assertEqual(first, D.tool_cache_dir(self.root, pdf))  # 稳定
        pdf.write_bytes(pdf.read_bytes())                          # 换个 mtime
        os.utime(pdf, (0, 0))
        self.assertNotEqual(first, D.tool_cache_dir(self.root, pdf))

    def test_unwritable_workspace_degrades_to_text_only(self):
        """工作空间不可写 → **正文照常** + 一句警告，绝不连文本一起丢。

        用"`.aigent` 是个文件"来模拟不可写：`mkdir(parents=True)` 必然失败，
        且这个构造在任何平台上都成立（不依赖 chmod 与运行用户）。
        """
        self.pdf("mixed.pdf", pages=("text", "image"))
        (self.root / D.TOOL_CACHE_DIRNAME).write_text("占位文件", encoding="utf-8")
        out = self.reg.run_read("mixed.pdf")
        self.assertIsInstance(out, str)          # 不是图片块
        self.assertFalse(out.startswith("Error:"), out)
        self.assertIn("quick brown fox", out)    # 文本层完好
        self.assertIn("页图目录不可写", out)      # 如实说明页图没给
        self.assertIn("本次未随附页图", out)
        self.assertNotIn("随附页图 1 张", out)     # 没有图片随附（不是"以为随附了"）

    def test_stale_cache_dirs_are_pruned(self):
        pdf = self.pdf("scan.pdf", pages=("image",))
        stale = self._cache_root() / "deadbeefdeadbeef"
        stale.mkdir(parents=True, exist_ok=True)
        (stale / "old.jpg").write_bytes(b"x")
        os.utime(stale, (0, 0))
        with mock.patch.dict("os.environ", {"TOOL_DOC_CACHE_TTL_SECONDS": "60"}):
            self.reg.run_read("scan.pdf")
        self.assertFalse(stale.exists(), "超期缓存目录应被清剪")


# ══════════════════════════════════════════════════════════════════
#  6. 多图聚合：一个组 = 一次文档读取
# ══════════════════════════════════════════════════════════════════

class ToolImageGroupTests(unittest.TestCase):
    def _pdf_like(self, pages: int, name: str = "spec.pdf") -> dict:
        images = [{"path": f"/w/p{i}.jpg", "name": f"{name} 第 {i} 页",
                   "mime": "image/jpeg", "page": i} for i in range(1, pages + 1)]
        return A.build_tool_images_result(images, text="正文", source=name)

    def test_one_result_may_carry_many_images(self):
        value = self._pdf_like(3)
        self.assertTrue(A.is_tool_image_result(value))
        self.assertEqual(len(A.tool_image_items(value)), 3)
        block = A.tool_image_block(value)
        self.assertEqual(len(block["images"]), 3)   # 一个结果 = 一块

    def test_group_is_admitted_whole_or_dropped_whole(self):
        """超上限时**整组**丢，绝不切页 —— 半份页图比没有更危险。"""
        msg = A.build_tool_images_message([self._pdf_like(3), self._pdf_like(2, "b.pdf")],
                                          limit=1)
        blocks = [b for b in msg["content"] if b["type"] == A.TOOL_IMAGE_BLOCK_TYPE]
        self.assertEqual(len(blocks), 1)
        self.assertEqual(len(blocks[0]["images"]), 3)      # 整组在
        text = msg["content"][0]["text"]
        self.assertIn("另有 1 个文件未随附", text)
        self.assertIn("b.pdf", text)

    def test_head_text_maps_pages_by_number(self):
        msg = A.build_tool_images_message([self._pdf_like(3)])
        head = msg["content"][0]["text"]
        self.assertIn("spec.pdf 第 1、2、3 页", head)
        self.assertIn("共 3 张", head)

    def test_legacy_singular_image_field_is_still_normalized(self):
        """存量 jsonl 里的单数 `image` 字段必须继续能读（回放不能瞎）。"""
        legacy = {"type": A.TOOL_IMAGE_BLOCK_TYPE, "text": "x",
                  "image": {"path": "/w/old.png", "name": "old.png",
                            "mime": "image/png"}}
        self.assertTrue(A.is_tool_image_result(legacy))
        self.assertEqual([i["path"] for i in A.tool_image_items(legacy)], ["/w/old.png"])
        self.assertEqual(len(A.tool_image_block(legacy)["images"]), 1)

    def test_multi_image_block_expands_to_many_image_urls(self):
        msg = A.build_tool_images_message([self._pdf_like(2)])
        out = A.expand_content_for_model(msg, None, supports_image=False)
        # 能力不符 → 每张图各降级成一句占位（位置保留，模型知道自己漏看了哪张）
        placeholders = [b for b in out["content"]
                        if b.get("type") == "text" and "未发送" in b.get("text", "")]
        self.assertEqual(len(placeholders), 2)

    def test_marker_is_not_in_model_message_whitelist(self):
        import agent_full_v2
        self.assertNotIn(A.TOOL_IMAGES_MARKER, agent_full_v2.MODEL_MSG_FIELDS)

    def test_no_bytes_anywhere_in_ledger(self):
        msg = A.build_tool_images_message([self._pdf_like(2)])
        blob = json.dumps(msg, ensure_ascii=False)
        self.assertNotIn("base64", blob)
        self.assertNotIn("data:image", blob)


# ══════════════════════════════════════════════════════════════════
#  7. 引用通道的注入文本与忽略清单
# ══════════════════════════════════════════════════════════════════

class RefsRoutingTests(unittest.TestCase):
    def test_injection_points_at_the_single_entry(self):
        text = R._render_ref_text([{"path": "/w/a.pdf", "is_dir": False}])
        self.assertIn("run_read", text)
        self.assertIn("自动分派", text)
        # 旧的路由表（哪个格式用哪个工具）必须消失 —— 那是"选错工具"的温床
        self.assertNotIn("run_read_pdf", text)
        self.assertNotIn("view_image", text)
        self.assertIn("不要在未读取的情况下猜测", text)

    def test_cache_dir_is_in_the_ignore_list(self):
        self.assertIn(D.TOOL_CACHE_DIRNAME, R.DEFAULT_IGNORE_DIRS)
        self.assertIn(D.TOOL_CACHE_DIRNAME, R.ignore_dirs())

    def test_cache_dir_never_shows_up_in_the_picker(self):
        """缓存目录不能出现在 @ 候选里：它既无价值，又能被选中引用。"""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp).resolve()
            (root / "keep.txt").write_text("x", encoding="utf-8")
            (root / D.TOOL_CACHE_DIRNAME / "pages" / "abc").mkdir(parents=True)
            out = R.list_workspace(root)
            names = {item["name"] for item in out["items"]}
            self.assertIn("keep.txt", names)
            self.assertNotIn(D.TOOL_CACHE_DIRNAME, names)


class SidePathImageAggregationTests(unittest.TestCase):
    """子智能体与队友这两条循环**没有独立的发送边界**，图片必须就地展开。

    违反后果（改造前真实存在）：`run_read` 读到图片返回 dict，被 `str(output)`
    成一段 Python repr 塞进 tool 消息 → 图彻底丢失，模型只看到一堆花括号。
    主智能体那条链路由 `test_view_image.py` 的链路级测试覆盖；这两条只能守源码
    （要真跑起来得驱动一整套子智能体循环），与仓库既有的源码守卫同款做法。
    """

    def _src(self, name: str) -> str:
        return (AGENTS_DIR / name).read_text(encoding="utf-8")

    def test_subagent_keeps_and_expands_image_blocks(self):
        src = self._src("subagent.py")
        self.assertIn("is_tool_image_result(output)", src)
        self.assertIn("tool_image_values.append(output)", src)
        self.assertIn("build_tool_images_message(tool_image_values)", src)
        self.assertIn("expand_content_for_model(", src)      # 就地展开成线格式
        self.assertIn("TOOL_IMAGES_MARKER, None", src)        # marker 必须摘掉
        self.assertNotIn('"content": str(output)', src)       # 不许再 str() 掉图片块
        # 追加点必须在 tool 消息循环**之后**（顺序硬约束）
        loop_at = src.index("for tool_call in sub_msg.tool_calls:")
        append_at = src.index("build_tool_images_message(tool_image_values)", loop_at)
        self.assertGreater(append_at, loop_at)

    def test_teammate_keeps_and_expands_image_blocks(self):
        src = self._src("teammate_manager.py")
        self.assertIn("is_tool_image_result(output)", src)
        self.assertIn("build_tool_images_message(tool_image_values)", src)
        self.assertIn("expand_content_for_model(", src)
        self.assertIn("TOOL_IMAGES_MARKER, None", src)
        loop_at = src.index("for tc in tool_calls:")
        append_at = src.index("build_tool_images_message(tool_image_values)", loop_at)
        self.assertGreater(append_at, loop_at)

    def test_teammate_tool_list_drops_retired_names(self):
        """队友的工具白名单也要跟着收敛，否则它仍然拿不到能读 PDF 的工具。"""
        src = self._src("teammate_manager.py")
        self.assertIn('wanted = {"bash", "run_read", "run_write"}', src)
        self.assertNotIn("run_read_pdf", src)
        self.assertNotIn("view_image", src)


if __name__ == "__main__":
    unittest.main()
